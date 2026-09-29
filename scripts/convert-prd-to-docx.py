#!/usr/bin/env python3
"""
将 Markdown PRD 文档转换为 DOCX 格式
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import re

def create_styled_document():
    """创建带样式的文档"""
    doc = Document()
    
    # 设置文档默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '微软雅黑'
    font.size = Pt(10.5)
    
    # 设置标题样式
    for i in range(1, 4):
        heading_style = doc.styles[f'Heading {i}']
        heading_style.font.name = '微软雅黑'
        heading_style.font.color.rgb = RGBColor(0, 0, 0)
        if i == 1:
            heading_style.font.size = Pt(18)
        elif i == 2:
            heading_style.font.size = Pt(15)
        else:
            heading_style.font.size = Pt(13)
    
    return doc

def parse_markdown_table(lines):
    """解析 Markdown 表格"""
    if len(lines) < 2:
        return None
    
    # 解析表头
    header = [cell.strip() for cell in lines[0].strip('|').split('|')]
    
    # 跳过分隔行
    data_lines = lines[2:] if len(lines) > 2 else []
    
    # 解析数据行
    rows = []
    for line in data_lines:
        if '|' in line and not line.strip().startswith('|---'):
            cells = [cell.strip() for cell in line.strip('|').split('|')]
            rows.append(cells)
    
    return header, rows

def add_table_to_doc(doc, header, rows):
    """添加表格到文档"""
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = 'Light Grid Accent 1'
    
    # 添加表头
    for i, cell_text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = cell_text
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(10)
    
    # 添加数据行
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            if j < len(table.rows[i+1].cells):
                cell = table.rows[i+1].cells[j]
                cell.text = cell_text
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
    
    doc.add_paragraph()  # 表格后加空行

def convert_markdown_to_docx(md_file, docx_file):
    """将 Markdown 转换为 DOCX"""
    doc = create_styled_document()
    
    with open(md_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        
        # 空行
        if not line:
            i += 1
            continue
        
        # 标题
        if line.startswith('#'):
            level = len(line.split()[0])
            text = line.lstrip('#').strip()
            doc.add_heading(text, level=min(level, 3))
            i += 1
            continue
        
        # 表格
        if '|' in line and i + 1 < len(lines) and '---' in lines[i + 1]:
            table_lines = []
            while i < len(lines) and '|' in lines[i]:
                table_lines.append(lines[i].rstrip())
                i += 1
            
            parsed = parse_markdown_table(table_lines)
            if parsed:
                header, rows = parsed
                add_table_to_doc(doc, header, rows)
            continue
        
        # 代码块
        if line.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i].rstrip())
                i += 1
            i += 1  # 跳过结束的 ```
            
            # 添加代码块（使用等宽字体）
            if code_lines:
                code_text = '\n'.join(code_lines)
                p = doc.add_paragraph()
                p.style = 'Normal'
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                p.paragraph_format.left_indent = Inches(0.3)
            continue
        
        # 列表项
        if line.startswith('- ') or line.startswith('* '):
            text = line[2:].strip()
            p = doc.add_paragraph(text, style='List Bullet')
            i += 1
            continue
        
        # 有序列表
        if re.match(r'^\d+\.\s', line):
            text = re.sub(r'^\d+\.\s', '', line).strip()
            p = doc.add_paragraph(text, style='List Number')
            i += 1
            continue
        
        # 普通段落
        if line.strip():
            # 处理加粗
            text = line.strip()
            p = doc.add_paragraph()
            
            # 简单的加粗处理
            parts = re.split(r'\*\*(.*?)\*\*', text)
            for j, part in enumerate(parts):
                if j % 2 == 1:  # 加粗部分
                    run = p.add_run(part)
                    run.font.bold = True
                else:
                    p.add_run(part)
        
        i += 1
    
    doc.save(docx_file)
    print(f"✅ 转换完成：{docx_file}")

if __name__ == '__main__':
    md_file = '/Users/smigoo/工作/mvgo/outputs/mvgo-workbench-prd.md'
    docx_file = '/Users/smigoo/工作/mvgo/outputs/mvgo-workbench-prd.docx'
    convert_markdown_to_docx(md_file, docx_file)
